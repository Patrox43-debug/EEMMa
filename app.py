"""
Backend API para la aplicación EEMM (Gestión de Equipos Médicos).
Construido con FastAPI y SQLite local.
"""

import os
import re
import io
import uuid
import base64
import sqlite3
import shutil
import tempfile
from datetime import datetime
from typing import Optional, List, Dict, Any
from PIL import Image, ImageOps

from fastapi import FastAPI, HTTPException, Query, Body, Response, UploadFile, File
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, JSONResponse
from fastapi.middleware.cors import CORSMiddleware
from starlette.background import BackgroundTask
import csv
from pydantic import BaseModel

from pdf_generator import generate_chequeo_pdf

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.getenv("DB_PATH", os.path.join(BASE_DIR, "eemm.db"))
STATIC_DIR = os.path.join(BASE_DIR, "static")
UPLOADS_DIR = os.getenv("UPLOADS_DIR", os.path.join(BASE_DIR, "uploads"))

os.makedirs(STATIC_DIR, exist_ok=True)
os.makedirs(UPLOADS_DIR, exist_ok=True)
if os.path.dirname(DB_PATH):
    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)

# Si se usa un volumen persistente y aún no existe la base de datos, copiar la base inicial completa
default_db = os.path.join(BASE_DIR, "eemm.db")
if os.path.abspath(DB_PATH) != os.path.abspath(default_db) and not os.path.exists(DB_PATH) and os.path.exists(default_db):
    try:
        import shutil
        shutil.copy2(default_db, DB_PATH)
        print(f"[INFO] Base de datos pre-cargada copiada exitosamente a {DB_PATH}")
    except Exception as e:
        print(f"[WARN] No se pudo copiar la base inicial a {DB_PATH}: {e}")

app = FastAPI(title="EEMM - Gestión de Equipos Médicos", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def ensure_db_initialized():
    """Asegura la creación de tablas e inicializa datos maestros si la base de datos está vacía."""
    try:
        conn = get_db()
        cursor = conn.cursor()

        # 1. Tabla perfiles
        cursor.execute("""
        CREATE TABLE IF NOT EXISTS perfiles (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nombre TEXT NOT NULL,
            rut TEXT UNIQUE,
            correo TEXT NOT NULL UNIQUE,
            password TEXT NOT NULL,
            funcion TEXT NOT NULL,
            tecnico TEXT NOT NULL,
            rol TEXT NOT NULL DEFAULT 'tecnico',
            created_at DATETIME DEFAULT CURRENT_TIMESTAMP
        );
        """)

        # 2. Tabla equipos
        cursor.execute("""
        CREATE TABLE IF NOT EXISTS equipos (
            id_equipo INTEGER PRIMARY KEY AUTOINCREMENT,
            codigo_origen INTEGER,
            nombre TEXT NOT NULL,
            marca TEXT,
            modelo TEXT,
            serie TEXT,
            categoria TEXT,
            estado TEXT DEFAULT 'Operativo',
            ubicacion TEXT,
            detalles TEXT,
            created_at DATETIME DEFAULT CURRENT_TIMESTAMP
        );
        """)
        cursor.execute("CREATE INDEX IF NOT EXISTS idx_equipos_nombre ON equipos(nombre);")
        cursor.execute("CREATE INDEX IF NOT EXISTS idx_equipos_serie ON equipos(serie);")
        cursor.execute("CREATE INDEX IF NOT EXISTS idx_equipos_categoria ON equipos(categoria);")

        # 3. Tabla registros
        cursor.execute("""
        CREATE TABLE IF NOT EXISTS registros (
            id_registro INTEGER PRIMARY KEY AUTOINCREMENT,
            fecha DATETIME NOT NULL,
            usuario TEXT NOT NULL,
            nombre_equipo TEXT NOT NULL,
            marca TEXT,
            modelo TEXT,
            serie TEXT,
            unidad TEXT,
            categoria TEXT,
            respuestas TEXT,
            obs TEXT,
            idpdf TEXT,
            foto_1 TEXT,
            foto_2 TEXT,
            foto_3 TEXT,
            foto_4 TEXT,
            firma_data TEXT,
            firma_nombre TEXT,
            created_at DATETIME DEFAULT CURRENT_TIMESTAMP
        );
        """)
        cursor.execute("CREATE INDEX IF NOT EXISTS idx_registros_fecha ON registros(fecha);")
        cursor.execute("CREATE INDEX IF NOT EXISTS idx_registros_usuario ON registros(usuario);")

        try:
            cursor.execute("ALTER TABLE registros ADD COLUMN categoria TEXT;")
        except Exception:
            pass

        # 4. Tabla reparaciones
        cursor.execute("""
        CREATE TABLE IF NOT EXISTS reparaciones (
            id_reparacion INTEGER PRIMARY KEY AUTOINCREMENT,
            fecha DATETIME,
            usuario TEXT,
            equipo TEXT,
            marca TEXT,
            modelo TEXT,
            serie TEXT,
            destino TEXT,
            parte_reparada TEXT,
            obs TEXT,
            created_at DATETIME DEFAULT CURRENT_TIMESTAMP
        );
        """)

        # 5. Tabla inventario
        cursor.execute("""
        CREATE TABLE IF NOT EXISTS inventario (
            id_caja TEXT PRIMARY KEY,
            nombre_caja TEXT NOT NULL,
            cantidad INTEGER DEFAULT 0,
            seccion_id TEXT NOT NULL,
            estanteria_id TEXT NOT NULL,
            barcode TEXT,
            finicio DATETIME,
            ftermino DATETIME,
            estado TEXT DEFAULT 'OK',
            created_at DATETIME DEFAULT CURRENT_TIMESTAMP
        );
        """)
        cursor.execute("CREATE INDEX IF NOT EXISTS idx_inventario_estanteria ON inventario(estanteria_id);")

        # 6. Tabla unidades
        cursor.execute("""
        CREATE TABLE IF NOT EXISTS unidades (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nombre TEXT NOT NULL UNIQUE
        );
        """)

        # Sembrar usuarios por defecto si está vacía
        cursor.execute("SELECT COUNT(*) FROM perfiles")
        if cursor.fetchone()[0] == 0:
            usuarios_iniciales = [
                ("Administrador General", "20.967.660-5", "admin@eemm.cl", "4277", "Jefe de Servicio", "Administrador del Sistema", "administrador"),
                ("Patricio Bustamante", "18.123.456-7", "patriciobustamante.ec@gmail.com", "1234", "Tecnico EEMM", "Patricio Bustamante", "tecnico"),
                ("Martin Peralta", "17.234.567-8", "martin.peralta@gmail.com", "1234", "Tecnico EEMM", "Martin Peralta", "tecnico"),
                ("Luis Vallejos", "16.345.678-9", "equiposmedicos1.@gmail.com", "1234", "Tecnico EEMM", "Luis Vallejos", "tecnico"),
                ("Jorge Ambrosetti", "15.456.789-0", "jambrosettic@gmail.com", "1234", "Tecnico EEMM", "Jorge Ambrosetti", "tecnico")
            ]
            cursor.executemany("""
                INSERT OR IGNORE INTO perfiles (nombre, rut, correo, password, funcion, tecnico, rol)
                VALUES (?, ?, ?, ?, ?, ?, ?);
            """, usuarios_iniciales)

        # Sembrar unidades si está vacía
        cursor.execute("SELECT COUNT(*) FROM unidades")
        if cursor.fetchone()[0] == 0:
            unidades_base = [
                "URGENCIA", "PABELLONES-QUIRURGICOS", "UPC", "ATENCION-OBSTETRICA",
                "GINECOLOGIA-Y-OBSTETRICIA", "LABORATORIO", "AISLAMIENTO",
                "MEDICO-QUIRURGICA-ADULTO", "MEDICO-QUIRURGICA-PEDIATRICA",
                "DIALISIS", "KINESIOTERAPIA-Y-REHABILITACION", "PARTOS",
                "PATOLOGICA(MORGUE)", "SOCIO-SANITARIO", "CIRUGIA-MENOR", "GES"
            ]
            cursor.executemany("INSERT OR IGNORE INTO unidades (nombre) VALUES (?);", [(u,) for u in unidades_base])

        # Sembrar equipos si está vacía y existe CSV
        cursor.execute("SELECT COUNT(*) FROM equipos")
        if cursor.fetchone()[0] == 0:
            equipos_csv = os.path.join(BASE_DIR, "csv", "equipos.csv")
            if os.path.exists(equipos_csv):
                with open(equipos_csv, mode="r", encoding="utf-8-sig") as f:
                    reader = csv.DictReader(f)
                    batch = []
                    for row in reader:
                        cod = row.get("codigo_origen")
                        cod_int = int(cod) if cod and cod.isdigit() else None
                        batch.append((
                            cod_int,
                            row.get("nombre", "SIN NOMBRE") or "SIN NOMBRE",
                            row.get("marca", "") or None,
                            row.get("modelo", "") or None,
                            row.get("serie", "") or None,
                            row.get("categoria", "") or None,
                            row.get("estado", "Operativo") or "Operativo",
                            row.get("ubicacion", "") or None,
                            row.get("detalles", "") or None
                        ))
                    cursor.executemany("""
                        INSERT INTO equipos (codigo_origen, nombre, marca, modelo, serie, categoria, estado, ubicacion, detalles)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?);
                    """, batch)

        # Sembrar inventario si está vacía y existe CSV
        cursor.execute("SELECT COUNT(*) FROM inventario")
        if cursor.fetchone()[0] == 0:
            inv_csv = os.path.join(BASE_DIR, "csv", "inventario.csv")
            if os.path.exists(inv_csv):
                with open(inv_csv, mode="r", encoding="utf-8-sig") as f:
                    reader = csv.DictReader(f)
                    batch = []
                    for row in reader:
                        cant = row.get("cantidad", "0")
                        cant_int = int(cant) if cant and cant.isdigit() else 0
                        batch.append((
                            row.get("id_caja"),
                            row.get("nombre_caja", "Caja"),
                            cant_int,
                            row.get("seccion_id", ""),
                            row.get("estanteria_id", ""),
                            row.get("barcode", ""),
                            row.get("finicio") or None,
                            row.get("ftermino") or None,
                            row.get("estado", "OK") or "OK"
                        ))
                    cursor.executemany("""
                        INSERT OR REPLACE INTO inventario (id_caja, nombre_caja, cantidad, seccion_id, estanteria_id, barcode, finicio, ftermino, estado)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?);
                    """, batch)

        conn.commit()
        conn.close()
    except Exception as e:
        print(f"[WARN] Error en ensure_db_initialized: {e}")

# Ejecutar inicialización garantizada de base de datos
ensure_db_initialized()

@app.get("/api/health")
def health_check():
    """Endpoint de salud para Coolify, Docker y balanceadores de carga."""
    return {
        "status": "healthy",
        "app": "EEMM",
        "database": os.path.exists(DB_PATH),
        "timestamp": datetime.now().isoformat()
    }

def save_base64_image(data_uri: str, prefix: str) -> Optional[str]:
    """Guarda una imagen enviada como base64 data URI a disco con compresión y retorna la ruta relativa."""
    if not data_uri or not isinstance(data_uri, str) or not data_uri.startswith("data:image"):
        return None
    try:
        header, encoded = data_uri.split(",", 1)
        raw_bytes = base64.b64decode(encoded)
        img = Image.open(io.BytesIO(raw_bytes))

        # Corregir orientación EXIF si proviene de dispositivos móviles
        try:
            img = ImageOps.exif_transpose(img)
        except Exception:
            pass

        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        uid = uuid.uuid4().hex[:6]

        if prefix.startswith("firma"):
            # Para firmas: redimensionar si supera 800px y comprimir como PNG optimizado
            if img.width > 800:
                new_height = max(1, int(img.height * (800 / img.width)))
                img = img.resize((800, new_height), Image.Resampling.LANCZOS)
            filename = f"{prefix}_{timestamp}_{uid}.png"
            filepath = os.path.join(UPLOADS_DIR, filename)
            img.save(filepath, format="PNG", optimize=True)
        else:
            # Para fotos: convertir a RGB (reemplazar canal alfa por fondo blanco), redimensionar y comprimir JPEG
            if img.mode in ("RGBA", "LA", "P"):
                background = Image.new("RGB", img.size, (255, 255, 255))
                if img.mode == "P":
                    img = img.convert("RGBA")
                background.paste(img, mask=img.split()[-1] if img.mode == "RGBA" else None)
                img = background
            elif img.mode != "RGB":
                img = img.convert("RGB")

            # Redimensionar a un máximo de 1280px para mantener nitidez reduciendo peso drásticamente
            max_dim = 1280
            if img.width > max_dim or img.height > max_dim:
                img.thumbnail((max_dim, max_dim), Image.Resampling.LANCZOS)

            filename = f"{prefix}_{timestamp}_{uid}.jpg"
            filepath = os.path.join(UPLOADS_DIR, filename)
            # Calidad 78 optimizada: reduce el tamaño a ~60-120 KB por foto
            img.save(filepath, format="JPEG", quality=78, optimize=True)

        return f"/uploads/{filename}"
    except Exception as e:
        print(f"Error procesando/comprimiendo imagen ({prefix}): {e}")
        return None

def normalize_rut(rut: str) -> str:
    """Normaliza un RUT chileno eliminando puntos, guiones y espacios."""
    if not rut:
        return ""
    return re.sub(r'[^0-9kK]', '', str(rut)).upper()

# Modelos Pydantic
class LoginRequest(BaseModel):
    rut: str
    password: str

class UserCreate(BaseModel):
    nombre: str
    rut: str
    correo: str
    password: str
    rol: Optional[str] = "tecnico"
    funcion: Optional[str] = "Tecnico EEMM"

class UserUpdate(BaseModel):
    nombre: str
    rut: str
    correo: str
    password: Optional[str] = None
    rol: Optional[str] = "tecnico"
    funcion: Optional[str] = "Tecnico EEMM"

class ChequeoCreate(BaseModel):
    usuario: str
    nombre_equipo: str
    marca: Optional[str] = ""
    modelo: Optional[str] = ""
    serie: Optional[str] = ""
    unidad: Optional[str] = ""
    categoria: Optional[str] = "GENERAL"
    respuestas: Optional[Any] = None
    obs: Optional[str] = ""
    fotos: Optional[List[str]] = []  # Lista de hasta 4 data URIs
    firma: Optional[str] = None      # Data URI del canvas
    firma_nombre: Optional[str] = ""
    fecha: Optional[str] = None      # Fecha original en caso de guardado offline

# Rutas de Autenticación con RUT
@app.post("/api/auth/login")
def login(creds: LoginRequest):
    rut_clean = normalize_rut(creds.rut)
    if not rut_clean:
        raise HTTPException(status_code=400, detail="Por favor ingresa un RUT válido")

    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT id, nombre, rut, correo, password, funcion, tecnico, rol
        FROM perfiles
    """)
    users = cursor.fetchall()
    conn.close()

    matched_user = None
    for u in users:
        if normalize_rut(u["rut"]) == rut_clean:
            matched_user = u
            break

    if not matched_user:
        raise HTTPException(status_code=401, detail="RUT no registrado en el sistema")
    
    # Validar contraseña
    if str(matched_user["password"]).strip() != str(creds.password).strip():
        raise HTTPException(status_code=401, detail="Contraseña incorrecta")

    return {
        "status": "ok",
        "user": {
            "id": matched_user["id"],
            "nombre": matched_user["nombre"],
            "rut": matched_user["rut"],
            "correo": matched_user["correo"],
            "funcion": matched_user["funcion"],
            "tecnico": matched_user["tecnico"] or matched_user["nombre"],
            "rol": matched_user["rol"]
        }
    }

# ==============================================================================
# Rutas de Administración de Usuarios (CRUD)
# ==============================================================================
@app.get("/api/admin/users")
def get_admin_users():
    """Listar todos los usuarios para el panel de administración."""
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT id, nombre, rut, correo, password, funcion, tecnico, rol, created_at
        FROM perfiles
        ORDER BY id ASC
    """)
    users = [dict(row) for row in cursor.fetchall()]
    conn.close()
    return users

@app.post("/api/admin/users")
def create_admin_user(data: UserCreate):
    """Crear un nuevo usuario."""
    rut_clean = normalize_rut(data.rut)
    if not rut_clean:
        raise HTTPException(status_code=400, detail="El RUT proporcionado es inválido")

    conn = get_db()
    cursor = conn.cursor()

    # Verificar si el RUT ya existe
    cursor.execute("SELECT id, rut FROM perfiles")
    for row in cursor.fetchall():
        if normalize_rut(row["rut"]) == rut_clean:
            conn.close()
            raise HTTPException(status_code=400, detail=f"Ya existe un usuario con el RUT {data.rut}")

    # Verificar correo
    cursor.execute("SELECT id FROM perfiles WHERE LOWER(correo) = LOWER(?)", (data.correo.strip(),))
    if cursor.fetchone():
        conn.close()
        raise HTTPException(status_code=400, detail=f"Ya existe un usuario con el correo {data.correo}")

    cursor.execute("""
        INSERT INTO perfiles (nombre, rut, correo, password, funcion, tecnico, rol)
        VALUES (?, ?, ?, ?, ?, ?, ?);
    """, (
        data.nombre.strip(),
        data.rut.strip(),
        data.correo.strip(),
        data.password.strip(),
        data.funcion.strip() if data.funcion else "Tecnico EEMM",
        data.nombre.strip(),
        data.rol.strip() if data.rol in ["administrador", "tecnico"] else "tecnico"
    ))
    new_id = cursor.lastrowid
    conn.commit()
    conn.close()

    return {"status": "ok", "id": new_id, "message": "Usuario creado exitosamente"}

@app.put("/api/admin/users/{user_id}")
def update_admin_user(user_id: int, data: UserUpdate):
    """Actualizar datos de un usuario existente."""
    rut_clean = normalize_rut(data.rut)
    if not rut_clean:
        raise HTTPException(status_code=400, detail="El RUT proporcionado es inválido")

    conn = get_db()
    cursor = conn.cursor()

    cursor.execute("SELECT id FROM perfiles WHERE id = ?", (user_id,))
    if not cursor.fetchone():
        conn.close()
        raise HTTPException(status_code=404, detail="Usuario no encontrado")

    # Validar que otro usuario no tenga el mismo RUT
    cursor.execute("SELECT id, rut FROM perfiles WHERE id != ?", (user_id,))
    for row in cursor.fetchall():
        if normalize_rut(row["rut"]) == rut_clean:
            conn.close()
            raise HTTPException(status_code=400, detail=f"El RUT {data.rut} ya pertenece a otro usuario")

    # Actualizar con o sin contraseña
    if data.password and data.password.strip():
        cursor.execute("""
            UPDATE perfiles
            SET nombre = ?, rut = ?, correo = ?, password = ?, rol = ?, funcion = ?, tecnico = ?
            WHERE id = ?;
        """, (
            data.nombre.strip(),
            data.rut.strip(),
            data.correo.strip(),
            data.password.strip(),
            data.rol.strip(),
            data.funcion.strip() if data.funcion else "Tecnico EEMM",
            data.nombre.strip(),
            user_id
        ))
    else:
        cursor.execute("""
            UPDATE perfiles
            SET nombre = ?, rut = ?, correo = ?, rol = ?, funcion = ?, tecnico = ?
            WHERE id = ?;
        """, (
            data.nombre.strip(),
            data.rut.strip(),
            data.correo.strip(),
            data.rol.strip(),
            data.funcion.strip() if data.funcion else "Tecnico EEMM",
            data.nombre.strip(),
            user_id
        ))

    conn.commit()
    conn.close()
    return {"status": "ok", "message": "Usuario actualizado exitosamente"}

@app.delete("/api/admin/users/{user_id}")
def delete_admin_user(user_id: int):
    """Eliminar un usuario (no permite eliminar al último administrador)."""
    conn = get_db()
    cursor = conn.cursor()

    cursor.execute("SELECT id, rol FROM perfiles WHERE id = ?", (user_id,))
    user = cursor.fetchone()
    if not user:
        conn.close()
        raise HTTPException(status_code=404, detail="Usuario no encontrado")

    if user["rol"] == "administrador":
        cursor.execute("SELECT COUNT(*) FROM perfiles WHERE rol = 'administrador'")
        admin_count = cursor.fetchone()[0]
        if admin_count <= 1:
            conn.close()
            raise HTTPException(status_code=400, detail="No se puede eliminar el único administrador del sistema")

    cursor.execute("DELETE FROM perfiles WHERE id = ?", (user_id,))
    conn.commit()
    conn.close()
    return {"status": "ok", "message": "Usuario eliminado exitosamente"}

# ==============================================================================
# Rutas de Respaldo y Mantenimiento de Base de Datos (Exportar / Importar)
# ==============================================================================
@app.get("/api/admin/database/export")
def export_database():
    """Exportar copia de seguridad completa de la base de datos SQLite (.db)."""
    if not os.path.exists(DB_PATH):
        raise HTTPException(status_code=404, detail="Archivo de base de datos no encontrado")

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"eemm_backup_{timestamp}.db"
    temp_dir = tempfile.gettempdir()
    temp_backup_path = os.path.join(temp_dir, f"export_{uuid.uuid4().hex[:8]}_{filename}")

    try:
        # Usar API de respaldo nativo de SQLite para copia consistente y atómica
        src_conn = sqlite3.connect(DB_PATH)
        dst_conn = sqlite3.connect(temp_backup_path)
        with dst_conn:
            src_conn.backup(dst_conn)
        dst_conn.close()
        src_conn.close()

        def cleanup():
            if os.path.exists(temp_backup_path):
                try:
                    os.remove(temp_backup_path)
                except Exception:
                    pass

        return FileResponse(
            path=temp_backup_path,
            filename=filename,
            media_type="application/x-sqlite3",
            background=BackgroundTask(cleanup)
        )
    except Exception as e:
        if os.path.exists(temp_backup_path):
            try:
                os.remove(temp_backup_path)
            except Exception:
                pass
        raise HTTPException(status_code=500, detail=f"Error generando respaldo de la base de datos: {str(e)}")


@app.post("/api/admin/database/import")
async def import_database(file: UploadFile = File(...)):
    """Importar y restaurar la base de datos a partir de un archivo .db o .sql."""
    filename = file.filename or ""
    ext = os.path.splitext(filename)[1].lower()
    if ext not in [".db", ".sqlite", ".sqlite3", ".sql"]:
        raise HTTPException(
            status_code=400,
            detail="Formato no soportado. Debe seleccionar un archivo de base de datos (.db, .sqlite, .sqlite3 o .sql)"
        )

    temp_dir = tempfile.gettempdir()
    temp_import_path = os.path.join(temp_dir, f"import_{uuid.uuid4().hex[:8]}{ext}")
    temp_verified_db = os.path.join(temp_dir, f"verified_{uuid.uuid4().hex[:8]}.db")

    try:
        # Guardar archivo subido en ruta temporal
        contents = await file.read()
        if len(contents) == 0:
            raise HTTPException(status_code=400, detail="El archivo subido está vacío")

        with open(temp_import_path, "wb") as f:
            f.write(contents)

        # Si es .sql, creamos una base SQLite temporal y ejecutamos el script
        if ext == ".sql":
            sql_text = contents.decode("utf-8-sig", errors="replace")
            test_conn = sqlite3.connect(temp_verified_db)
            try:
                test_conn.executescript(sql_text)
                test_conn.commit()
            except Exception as sql_err:
                test_conn.close()
                raise HTTPException(
                    status_code=400,
                    detail=f"Error al procesar archivo SQL: {str(sql_err)}"
                )
        else:
            # Es un archivo SQLite binario (.db)
            # Verificar cabecera mágica de SQLite
            if len(contents) < 16 or contents[:15] != b"SQLite format 3":
                raise HTTPException(
                    status_code=400,
                    detail="El archivo proporcionado no es una base de datos SQLite válida (cabecera corrupta)"
                )
            shutil.copyfile(temp_import_path, temp_verified_db)
            test_conn = sqlite3.connect(temp_verified_db)

        # Verificar integridad y presencia de tablas requeridas
        cursor = test_conn.cursor()
        cursor.execute("PRAGMA quick_check;")
        check_res = cursor.fetchone()
        if not check_res or check_res[0] != "ok":
            test_conn.close()
            raise HTTPException(status_code=400, detail="La base de datos subida no superó la verificación de integridad de SQLite")

        cursor.execute("SELECT name FROM sqlite_master WHERE type='table';")
        tables = [row[0] for row in cursor.fetchall()]
        if "perfiles" not in tables:
            test_conn.close()
            raise HTTPException(
                status_code=400,
                detail="La base de datos no contiene la tabla esencial 'perfiles'. No es un respaldo válido de EEMM."
            )

        # Obtener estadísticas de los datos que se van a restaurar
        stats = {}
        for tbl in ["perfiles", "equipos", "registros", "inventario", "unidades", "reparaciones"]:
            if tbl in tables:
                cursor.execute(f"SELECT COUNT(*) FROM {tbl}")
                stats[tbl] = cursor.fetchone()[0]
            else:
                stats[tbl] = 0

        test_conn.close()

        # Generar respaldo automático de seguridad de la base actual antes de sobrescribir
        if os.path.exists(DB_PATH):
            backup_safety = f"{DB_PATH}.bak_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
            try:
                shutil.copyfile(DB_PATH, backup_safety)
                print(f"[INFO] Respaldo automático de seguridad creado en: {backup_safety}")
            except Exception as bak_err:
                print(f"[WARN] No se pudo crear respaldo preventivo: {bak_err}")

        # Reemplazar la base de datos activa con la verificada
        shutil.copyfile(temp_verified_db, DB_PATH)

        # Asegurar índices y estructuras actualizadas
        ensure_db_initialized()

        return {
            "status": "ok",
            "message": "Base de datos restaurada exitosamente",
            "filename": filename,
            "stats": {
                "usuarios": stats.get("perfiles", 0),
                "equipos": stats.get("equipos", 0),
                "chequeos": stats.get("registros", 0),
                "inventario": stats.get("inventario", 0),
                "unidades": stats.get("unidades", 0),
                "reparaciones": stats.get("reparaciones", 0)
            }
        }

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error al restaurar la base de datos: {str(e)}")
    finally:
        # Limpieza de temporales
        for p in [temp_import_path, temp_verified_db]:
            if os.path.exists(p):
                try:
                    os.remove(p)
                except Exception:
                    pass

# Rutas de Catálogos (Equipos y Unidades)
@app.get("/api/equipos")
def get_equipos(q: Optional[str] = Query(None, description="Término de búsqueda"), limit: int = 50):
    conn = get_db()
    cursor = conn.cursor()
    if q and q.strip():
        term = f"%{q.strip()}%"
        cursor.execute("""
            SELECT id_equipo, codigo_origen, nombre, marca, modelo, serie, categoria, estado, ubicacion, detalles
            FROM equipos
            WHERE nombre LIKE ? OR serie LIKE ? OR marca LIKE ? OR modelo LIKE ?
            ORDER BY nombre ASC
            LIMIT ?
        """, (term, term, term, term, limit))
    else:
        cursor.execute("""
            SELECT id_equipo, codigo_origen, nombre, marca, modelo, serie, categoria, estado, ubicacion, detalles
            FROM equipos
            ORDER BY id_equipo ASC
            LIMIT ?
        """, (limit,))
    
    equipos = [dict(row) for row in cursor.fetchall()]
    conn.close()
    return equipos

@app.get("/api/unidades")
def get_unidades():
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT id, nombre FROM unidades ORDER BY nombre ASC")
    unidades = [row["nombre"] for row in cursor.fetchall()]
    conn.close()
    return unidades

# Rutas de Chequeos Preventivos
@app.post("/api/chequeos")
def create_chequeo(data: ChequeoCreate):
    conn = get_db()
    cursor = conn.cursor()

    # Procesar hasta 4 fotos
    foto_rutas = [None, None, None, None]
    if data.fotos:
        for idx, foto_b64 in enumerate(data.fotos[:4]):
            if foto_b64:
                ruta = save_base64_image(foto_b64, f"foto_{idx+1}")
                foto_rutas[idx] = ruta

    # Procesar firma
    firma_ruta = None
    if data.firma:
        firma_ruta = save_base64_image(data.firma, "firma")

    # Formatear respuestas
    respuestas_str = ""
    if isinstance(data.respuestas, list):
        respuestas_str = " | ".join([f"{item.get('item', '')}: {item.get('val', 'OK')}" for item in data.respuestas])
    elif isinstance(data.respuestas, str):
        respuestas_str = data.respuestas

    fecha_efectiva = (data.fecha and data.fecha.strip()) or datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    cursor.execute("""
        INSERT INTO registros (
            fecha, usuario, nombre_equipo, marca, modelo, serie, unidad, categoria,
            respuestas, obs, idpdf, foto_1, foto_2, foto_3, foto_4, firma_data, firma_nombre
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
    """, (
        fecha_efectiva,
        data.usuario or "Técnico",
        data.nombre_equipo or "Equipo",
        data.marca or "",
        data.modelo or "",
        data.serie or "",
        data.unidad or "",
        (data.categoria or "GENERAL").strip().upper(),
        respuestas_str,
        data.obs or "",
        f"PDF-{int(datetime.now().timestamp())}",
        foto_rutas[0],
        foto_rutas[1],
        foto_rutas[2],
        foto_rutas[3],
        firma_ruta,
        data.firma_nombre or ""
    ))
    new_id = cursor.lastrowid
    conn.commit()
    conn.close()

    return {
        "status": "ok",
        "id_registro": new_id,
        "message": "Chequeo preventivo registrado exitosamente",
        "fotos": [f for f in foto_rutas if f],
        "firma": firma_ruta
    }

@app.get("/api/chequeos")
def get_chequeos(
    usuario: Optional[str] = None,
    unidad: Optional[str] = None,
    q: Optional[str] = None,
    limit: int = 50
):
    conn = get_db()
    cursor = conn.cursor()

    conditions = []
    params = []

    if usuario and usuario.strip():
        conditions.append("LOWER(usuario) = LOWER(?)")
        params.append(usuario.strip())
    if unidad and unidad.strip():
        conditions.append("LOWER(unidad) = LOWER(?)")
        params.append(unidad.strip())
    if q and q.strip():
        term = f"%{q.strip()}%"
        conditions.append("(nombre_equipo LIKE ? OR serie LIKE ? OR obs LIKE ?)")
        params.extend([term, term, term])

    where_clause = "WHERE " + " AND ".join(conditions) if conditions else ""
    query = f"""
        SELECT id_registro, fecha, usuario, nombre_equipo, marca, modelo, serie,
               unidad, categoria, respuestas, obs, idpdf, foto_1, foto_2, foto_3, foto_4,
               firma_data, firma_nombre, created_at
        FROM registros
        {where_clause}
        ORDER BY id_registro DESC
        LIMIT ?
    """
    params.append(limit)
    cursor.execute(query, params)
    chequeos = [dict(row) for row in cursor.fetchall()]
    conn.close()
    return chequeos

@app.get("/api/chequeos/{id_reg}")
def get_chequeo_detail(id_reg: int):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT id_registro, fecha, usuario, nombre_equipo, marca, modelo, serie,
               unidad, categoria, respuestas, obs, idpdf, foto_1, foto_2, foto_3, foto_4,
               firma_data, firma_nombre, created_at
        FROM registros
        WHERE id_registro = ?
    """, (id_reg,))
    row = cursor.fetchone()
    conn.close()
    if not row:
        raise HTTPException(status_code=404, detail="Registro no encontrado")
    return dict(row)

@app.get("/api/chequeos/{id_reg}/pdf")
def get_chequeo_pdf(id_reg: int):
    """Genera y descarga el reporte oficial en PDF del chequeo preventivo bajo demanda."""
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT id_registro, fecha, usuario, nombre_equipo, marca, modelo, serie,
               unidad, categoria, respuestas, obs, idpdf, foto_1, foto_2, foto_3, foto_4,
               firma_data, firma_nombre, created_at
        FROM registros
        WHERE id_registro = ?
    """, (id_reg,))
    row = cursor.fetchone()
    conn.close()

    if not row:
        raise HTTPException(status_code=404, detail="Chequeo no encontrado")

    try:
        record = dict(row)
        pdf_bytes = generate_chequeo_pdf(record)
        filename = f"Chequeo_EEMM_{record.get('id_registro')}_{record.get('serie') or 'SN'}.pdf"
        return Response(
            content=pdf_bytes,
            media_type="application/pdf",
            headers={
                "Content-Disposition": f"inline; filename=\"{filename}\"",
                "Cache-Control": "no-cache"
            }
        )
    except Exception as e:
        print(f"Error generando PDF para chequeo #{id_reg}: {e}")
        raise HTTPException(status_code=500, detail=f"Error al generar PDF: {str(e)}")

# Métricas para Panel de Administrador
@app.get("/api/stats")
def get_stats():
    conn = get_db()
    cursor = conn.cursor()

    # Total equipos
    cursor.execute("SELECT COUNT(*) FROM equipos")
    total_equipos = cursor.fetchone()[0]

    # Total chequeos
    cursor.execute("SELECT COUNT(*) FROM registros")
    total_chequeos = cursor.fetchone()[0]

    # Chequeos con fotos o firmas
    cursor.execute("SELECT COUNT(*) FROM registros WHERE foto_1 IS NOT NULL OR firma_data IS NOT NULL")
    chequeos_con_multimedia = cursor.fetchone()[0]

    # Total inventario
    cursor.execute("SELECT COUNT(*), SUM(cantidad) FROM inventario")
    inv_row = cursor.fetchone()
    total_cajas = inv_row[0]
    total_unidades_stock = inv_row[1] or 0

    # Top unidades con más chequeos
    cursor.execute("""
        SELECT unidad, COUNT(*) as cant
        FROM registros
        WHERE unidad IS NOT NULL AND TRIM(unidad) != ''
        GROUP BY unidad
        ORDER BY cant DESC
        LIMIT 6
    """)
    top_unidades = [dict(row) for row in cursor.fetchall()]

    # Chequeos por técnico
    cursor.execute("""
        SELECT usuario, COUNT(*) as cant
        FROM registros
        GROUP BY usuario
        ORDER BY cant DESC
    """)
    por_tecnico = [dict(row) for row in cursor.fetchall()]

    conn.close()

    return {
        "total_equipos": total_equipos,
        "total_chequeos": total_chequeos,
        "chequeos_con_multimedia": chequeos_con_multimedia,
        "total_cajas_inventario": total_cajas,
        "total_unidades_stock": total_unidades_stock,
        "top_unidades": top_unidades,
        "por_tecnico": por_tecnico
    }

@app.get("/api/inventario")
def get_inventario(
    q: Optional[str] = None, 
    estanteria: Optional[str] = None,
    seccion: Optional[str] = None,
    estado: Optional[str] = None,
    limit: int = 500
):
    conn = get_db()
    cursor = conn.cursor()
    
    conditions = []
    params = []

    if q and q.strip():
        term = f"%{q.strip()}%"
        conditions.append("(nombre_caja LIKE ? OR id_caja LIKE ? OR barcode LIKE ? OR estanteria_id LIKE ?)")
        params.extend([term, term, term, term])

    if estanteria and estanteria.strip():
        conditions.append("estanteria_id = ?")
        params.append(estanteria.strip())

    if seccion and seccion.strip():
        conditions.append("seccion_id = ?")
        params.append(seccion.strip())

    if estado and estado.strip():
        conditions.append("estado = ?")
        params.append(estado.strip())

    sql = "SELECT id_caja, nombre_caja, cantidad, seccion_id, estanteria_id, barcode, finicio, ftermino, estado FROM inventario"
    if conditions:
        sql += " WHERE " + " AND ".join(conditions)
    sql += " ORDER BY estanteria_id ASC, seccion_id ASC, id_caja ASC LIMIT ?"
    params.append(limit)

    cursor.execute(sql, tuple(params))
    cajas = [dict(row) for row in cursor.fetchall()]
    conn.close()
    return cajas

@app.get("/api/inventario/export/excel")
def export_inventario_excel(
    q: Optional[str] = None,
    estado: Optional[str] = None,
    estanteria: Optional[str] = None
):
    """Exportar inventario de bodega en formato Excel profesional (.xlsx)."""
    conn = get_db()
    cursor = conn.cursor()

    # Mapeo de identificadores de estantes a números amigables
    cursor.execute("""
        SELECT DISTINCT estanteria_id
        FROM inventario
        WHERE estanteria_id IS NOT NULL AND estanteria_id != ''
        ORDER BY estanteria_id ASC
    """)
    shelf_rows = cursor.fetchall()
    shelf_map = {}
    for idx, sr in enumerate(shelf_rows, start=1):
        shelf_map[sr["estanteria_id"]] = f"Estante #{idx}"

    # Construir filtros opcionales
    conditions = []
    params = []
    
    if q and q.strip():
        term = f"%{q.strip()}%"
        conditions.append("(nombre_caja LIKE ? OR barcode LIKE ? OR id_caja LIKE ?)")
        params.extend([term, term, term])
        
    if estado and estado.strip():
        conditions.append("estado = ?")
        params.append(estado.strip())
        
    if estanteria and estanteria.strip():
        conditions.append("estanteria_id = ?")
        params.append(estanteria.strip())

    sql = """
        SELECT id_caja, nombre_caja, cantidad, seccion_id, estanteria_id, barcode, finicio, ftermino, estado
        FROM inventario
    """
    if conditions:
        sql += " WHERE " + " AND ".join(conditions)
    sql += " ORDER BY estanteria_id ASC, seccion_id ASC, nombre_caja ASC;"

    cursor.execute(sql, params)
    items = cursor.fetchall()
    conn.close()

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Inventario Bodega"
        ws.views.sheetView[0].showGridLines = True

        # Paleta y Estilos Visuales
        font_title = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")
        font_sub = Font(name="Segoe UI", size=9, italic=True, color="64748B")
        font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        fill_header = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
        font_data = Font(name="Segoe UI", size=10)
        font_bold = Font(name="Segoe UI", size=10, bold=True)
        fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        fill_ok = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        font_ok = Font(name="Segoe UI", size=9, bold=True, color="166534")

        fill_bajo = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
        font_bajo = Font(name="Segoe UI", size=9, bold=True, color="92400E")

        thin_side = Side(style="thin", color="CBD5E1")
        border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        thick_top = Border(top=Side(style="medium", color="1E40AF"), bottom=Side(style="double", color="1E40AF"))

        # Encabezado corporativo
        ws.merge_cells("A1:J1")
        ws["A1"] = "HOSPITAL CLÍNICO - REPORTE GENERAL DE BODEGA E INSUMOS"
        ws["A1"].font = font_title
        ws["A1"].alignment = Alignment(vertical="center")

        ws.merge_cells("A2:J2")
        ws["A2"] = f"Generado el: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')} | Total de cajas listadas: {len(items)}"
        ws["A2"].font = font_sub
        ws["A2"].alignment = Alignment(vertical="center")

        ws.row_dimensions[1].height = 24
        ws.row_dimensions[2].height = 18
        ws.row_dimensions[3].height = 6

        # Fila 4: Columnas
        headers = [
            ("ID Caja", 18),
            ("Descripción / Repuesto", 46),
            ("Stock (Un.)", 14),
            ("Estantería", 22),
            ("Nivel / Repisa", 24),
            ("Código Barra", 20),
            ("Estado", 16),
            ("Código Sección", 22),
            ("Fecha Registro", 20),
            ("Fecha Vencimiento", 20),
        ]

        header_row = 4
        ws.row_dimensions[header_row].height = 26

        for col_idx, (col_title, col_width) in enumerate(headers, start=1):
            cell = ws.cell(row=header_row, column=col_idx, value=col_title)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border_cell

        # Filas de datos
        total_unidades = 0
        current_row = 5

        for item in items:
            sec_id = item["seccion_id"] or ""
            repisa_num = 1
            if "_sec_" in sec_id:
                try:
                    repisa_num = int(sec_id.split("_sec_")[-1])
                except ValueError:
                    repisa_num = 1
            
            repisa_txt = f"Repisa {repisa_num}"
            if repisa_num == 1:
                repisa_txt += " (Superior)"
            elif repisa_num == 2:
                repisa_txt += " (Medio)"
            elif repisa_num == 3:
                repisa_txt += " (Inferior)"
            elif repisa_num == 4:
                repisa_txt += " (Base)"

            estante_txt = shelf_map.get(item["estanteria_id"], item["estanteria_id"] or "S/E")
            cant = int(item["cantidad"] or 0)
            total_unidades += cant
            is_bajo = (item["estado"] != "OK")

            ws.row_dimensions[current_row].height = 20
            is_even = (current_row % 2 == 0)
            row_fill = fill_zebra if is_even else PatternFill(fill_type=None)

            c1 = ws.cell(row=current_row, column=1, value=item["id_caja"])
            c1.alignment = Alignment(horizontal="center", vertical="center")

            c2 = ws.cell(row=current_row, column=2, value=item["nombre_caja"])
            c2.alignment = Alignment(horizontal="left", vertical="center")

            c3 = ws.cell(row=current_row, column=3, value=cant)
            c3.alignment = Alignment(horizontal="right", vertical="center")
            c3.number_format = "#,##0"

            c4 = ws.cell(row=current_row, column=4, value=estante_txt)
            c4.alignment = Alignment(horizontal="left", vertical="center")

            c5 = ws.cell(row=current_row, column=5, value=repisa_txt)
            c5.alignment = Alignment(horizontal="left", vertical="center")

            c6 = ws.cell(row=current_row, column=6, value=item["barcode"] or "S/C")
            c6.alignment = Alignment(horizontal="center", vertical="center")

            c7 = ws.cell(row=current_row, column=7, value="STOCK BAJO" if is_bajo else "OK")
            c7.alignment = Alignment(horizontal="center", vertical="center")
            c7.font = font_bajo if is_bajo else font_ok
            c7.fill = fill_bajo if is_bajo else fill_ok

            c8 = ws.cell(row=current_row, column=8, value=sec_id)
            c8.alignment = Alignment(horizontal="center", vertical="center")

            c9 = ws.cell(row=current_row, column=9, value=item["finicio"] or "-")
            c9.alignment = Alignment(horizontal="center", vertical="center")

            c10 = ws.cell(row=current_row, column=10, value=item["ftermino"] or "-")
            c10.alignment = Alignment(horizontal="center", vertical="center")

            for col_idx in range(1, 11):
                cell = ws.cell(row=current_row, column=col_idx)
                if col_idx != 7:
                    cell.font = font_data
                    if row_fill.fill_type:
                        cell.fill = row_fill
                cell.border = border_cell

            current_row += 1

        # Fila de Totales
        ws.row_dimensions[current_row].height = 24
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=2)
        total_label = ws.cell(row=current_row, column=1, value=f"TOTALES ({len(items)} cajas registradas)")
        total_label.font = font_bold
        total_label.alignment = Alignment(horizontal="right", vertical="center")
        total_label.border = thick_top

        c_tot = ws.cell(row=current_row, column=3, value=total_unidades)
        c_tot.font = font_bold
        c_tot.alignment = Alignment(horizontal="right", vertical="center")
        c_tot.number_format = "#,##0"
        c_tot.border = thick_top

        for c in range(4, 11):
            cell = ws.cell(row=current_row, column=c)
            cell.border = thick_top

        for col_idx, (_, col_width) in enumerate(headers, start=1):
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = col_width

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        filename = f"Inventario_Bodega_EEMM_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
        return Response(
            content=output.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f"attachment; filename=\"{filename}\"",
                "Cache-Control": "no-cache"
            }
        )

    except Exception as e:
        print("Error generando Excel con openpyxl, generando CSV fallback:", e)
        output = io.StringIO()
        output.write('\ufeff')
        writer = csv.writer(output, delimiter=';', quoting=csv.QUOTE_MINIMAL)
        writer.writerow(["ID Caja", "Descripción / Repuesto", "Stock (Un.)", "Estantería", "Nivel / Repisa", "Código Barra", "Estado", "Código Sección", "Fecha Inicio", "Fecha Término"])
        for it in items:
            sec_id = it["seccion_id"] or ""
            rep_n = sec_id.split("_sec_")[-1] if "_sec_" in sec_id else "1"
            writer.writerow([
                it["id_caja"],
                it["nombre_caja"],
                it["cantidad"],
                shelf_map.get(it["estanteria_id"], it["estanteria_id"]),
                f"Repisa {rep_n}",
                it["barcode"] or "S/C",
                it["estado"],
                sec_id,
                it["finicio"] or "",
                it["ftermino"] or ""
            ])
        filename = f"Inventario_Bodega_EEMM_{datetime.now().strftime('%Y%m%d_%H%M')}.csv"
        return Response(
            content=output.getvalue().encode('utf-8-sig'),
            media_type="text/csv",
            headers={
                "Content-Disposition": f"attachment; filename=\"{filename}\"",
                "Cache-Control": "no-cache"
            }
        )

@app.get("/api/inventario/estanterias")
def get_estanterias_summary():
    """Retorna los 16 estantes con sus métricas y repisas organizadas para la vista física."""
    conn = get_db()
    cursor = conn.cursor()

    cursor.execute("""
        SELECT 
            estanteria_id, 
            seccion_id,
            COUNT(*) as total_cajas,
            SUM(cantidad) as total_stock,
            SUM(CASE WHEN estado != 'OK' THEN 1 ELSE 0 END) as stock_bajo
        FROM inventario
        GROUP BY estanteria_id, seccion_id
        ORDER BY estanteria_id ASC, seccion_id ASC
    """)
    rows = cursor.fetchall()
    conn.close()

    estantes_dict = {}
    idx = 1
    for r in rows:
        eid = r["estanteria_id"]
        if eid not in estantes_dict:
            estantes_dict[eid] = {
                "id_estanteria": eid,
                "numero": idx,
                "nombre": f"Estante #{idx}",
                "total_cajas": 0,
                "total_stock": 0,
                "stock_bajo": 0,
                "repisas": []
            }
            idx += 1

        sec_id = r["seccion_id"]
        # Extraer número de sección: p.ej. "shelf_..._sec_1" -> 1
        repisa_num = 1
        if "_sec_" in sec_id:
            try:
                repisa_num = int(sec_id.split("_sec_")[-1])
            except ValueError:
                repisa_num = 1

        repisa_nombre = f"Repisa {repisa_num}"
        if repisa_num == 1:
            repisa_nombre += " (Nivel Superior)"
        elif repisa_num == 2:
            repisa_nombre += " (Nivel Medio)"
        elif repisa_num == 3:
            repisa_nombre += " (Nivel Inferior)"
        elif repisa_num == 4:
            repisa_nombre += " (Nivel Base)"

        estantes_dict[eid]["total_cajas"] += r["total_cajas"]
        estantes_dict[eid]["total_stock"] += (r["total_stock"] or 0)
        estantes_dict[eid]["stock_bajo"] += (r["stock_bajo"] or 0)
        estantes_dict[eid]["repisas"].append({
            "seccion_id": sec_id,
            "repisa_num": repisa_num,
            "nombre": repisa_nombre,
            "cajas_count": r["total_cajas"],
            "stock_count": r["total_stock"] or 0
        })

    return list(estantes_dict.values())

class InsumoCreate(BaseModel):
    id_caja: Optional[str] = None
    nombre_caja: str
    cantidad: int = 0
    estanteria_id: str
    seccion_id: str
    barcode: Optional[str] = ""
    finicio: Optional[str] = None
    ftermino: Optional[str] = None
    estado: Optional[str] = "OK"

class InsumoUpdate(BaseModel):
    nombre_caja: Optional[str] = None
    cantidad: Optional[int] = None
    estanteria_id: Optional[str] = None
    seccion_id: Optional[str] = None
    barcode: Optional[str] = None
    finicio: Optional[str] = None
    ftermino: Optional[str] = None
    estado: Optional[str] = None

class InsumoIngreso(BaseModel):
    cantidad: int
    estado: Optional[str] = None

@app.post("/api/inventario")
def create_insumo(item: InsumoCreate):
    conn = get_db()
    cursor = conn.cursor()
    
    # Generar id_caja si no se proporciona
    cid = item.id_caja.strip() if item.id_caja and item.id_caja.strip() else f"box_{int(datetime.now().timestamp() * 1000)}"
    
    cursor.execute("SELECT id_caja FROM inventario WHERE id_caja = ?", (cid,))
    if cursor.fetchone():
        conn.close()
        raise HTTPException(status_code=400, detail=f"Ya existe un insumo con el código {cid}")

    cursor.execute("""
        INSERT INTO inventario (id_caja, nombre_caja, cantidad, seccion_id, estanteria_id, barcode, finicio, ftermino, estado)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        cid,
        item.nombre_caja.strip(),
        max(0, item.cantidad),
        item.seccion_id.strip(),
        item.estanteria_id.strip(),
        (item.barcode or "").strip(),
        item.finicio or datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        item.ftermino,
        item.estado or "OK"
    ))
    conn.commit()
    
    cursor.execute("SELECT * FROM inventario WHERE id_caja = ?", (cid,))
    row = cursor.fetchone()
    conn.close()
    return {"ok": True, "caja": dict(row)}

@app.put("/api/inventario/{id_caja}")
def update_insumo(id_caja: str, item: InsumoUpdate):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM inventario WHERE id_caja = ?", (id_caja,))
    curr = cursor.fetchone()
    if not curr:
        conn.close()
        raise HTTPException(status_code=404, detail="Insumo no encontrado")
    
    curr = dict(curr)
    nombre = item.nombre_caja.strip() if item.nombre_caja is not None else curr["nombre_caja"]
    cantidad = item.cantidad if item.cantidad is not None else curr["cantidad"]
    seccion = item.seccion_id.strip() if item.seccion_id is not None else curr["seccion_id"]
    estanteria = item.estanteria_id.strip() if item.estanteria_id is not None else curr["estanteria_id"]
    barcode = item.barcode.strip() if item.barcode is not None else curr["barcode"]
    estado = item.estado.strip() if item.estado is not None else curr["estado"]
    
    cursor.execute("""
        UPDATE inventario
        SET nombre_caja = ?, cantidad = ?, seccion_id = ?, estanteria_id = ?, barcode = ?, estado = ?
        WHERE id_caja = ?
    """, (nombre, max(0, cantidad), seccion, estanteria, barcode, estado, id_caja))
    conn.commit()
    
    cursor.execute("SELECT * FROM inventario WHERE id_caja = ?", (id_caja,))
    updated = cursor.fetchone()
    conn.close()
    return {"ok": True, "caja": dict(updated)}

@app.post("/api/inventario/{id_caja}/ingreso")
def ingresar_stock(id_caja: str, ingreso: InsumoIngreso):
    if ingreso.cantidad <= 0:
        raise HTTPException(status_code=400, detail="La cantidad a ingresar debe ser mayor a 0")
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM inventario WHERE id_caja = ?", (id_caja,))
    curr = cursor.fetchone()
    if not curr:
        conn.close()
        raise HTTPException(status_code=404, detail="Insumo no encontrado")
    
    curr_cant = curr["cantidad"] or 0
    nueva_cant = curr_cant + ingreso.cantidad
    nuevo_estado = ingreso.estado or ("OK" if nueva_cant >= 5 else curr["estado"])
    
    cursor.execute("""
        UPDATE inventario
        SET cantidad = ?, estado = ?
        WHERE id_caja = ?
    """, (nueva_cant, nuevo_estado, id_caja))
    conn.commit()
    
    cursor.execute("SELECT * FROM inventario WHERE id_caja = ?", (id_caja,))
    updated = cursor.fetchone()
    conn.close()
    return {"ok": True, "caja": dict(updated)}

@app.delete("/api/inventario/{id_caja}")
def delete_insumo(id_caja: str):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT id_caja, nombre_caja FROM inventario WHERE id_caja = ?", (id_caja,))
    curr = cursor.fetchone()
    if not curr:
        conn.close()
        raise HTTPException(status_code=404, detail="Insumo no encontrado")
    
    cursor.execute("DELETE FROM inventario WHERE id_caja = ?", (id_caja,))
    conn.commit()
    conn.close()
    return {"ok": True, "message": f"Insumo '{curr['nombre_caja']}' eliminado correctamente"}


# Catálogo oficial de pautas técnicas por categoría
PAUTAS_POR_CATEGORIA = {
    "GENERAL": [
        "Chasis / Carcasa", "Ruedas / Frenos", "Cables alimentación", "Interruptores / Fusibles", 
        "Tubos / Mangueras", "Transductores", "Controles / Perillas", "Indicadores / Display", 
        "Señales audibles", "Control remoto / Pedal", "Batería", "Accesorios", 
        "Limpieza exterior", "Limpieza interior", "Lubricación", "Calibración", 
        "Reemplazo partes", "Chequeo fugas", "Revisión presiones", "Vaporizadores"
    ],
    "VENTILACION": [
        "Estado de mangueras de oxígeno/aire", "Válvula exhalatoria", "Celda de O2", 
        "Prueba de fugas de circuito", "Filtros antibacterianos", "Sensores de flujo", 
        "Verificación de alarmas", "Batería interna (autonomía)", "Compresor (si aplica)", 
        "Humidificador", "Limpieza de filtros de aire"
    ],
    "MONITOREO": [
        "Chasis / Carcasa", "Limpieza exterior", "Cables alimentación", "Alarmas", 
        "Estado Bateria", "Manguera PANI", "Manguito PANI", "Cable SPO2", 
        "Señales audibles", "Estado Pantalla/Display", "Estado Botones Interruptores", "Accesorios", 
        "Limpieza interior", "Calibración", "Cable ECG", "Estado de Palas", "Estado de Parches"
    ],
    "ASPIRACION": [
        "Chasis / Carcasa", "Limpieza exterior", "Cables alimentación", 
        "Estado Bateria", "Estado Botones Interruptores", "Accesorios"
    ],
    "LABORATORIO": [
        "Chasis / Carcasa", "Limpieza exterior", "Cables alimentación", "Alarmas", 
        "Estado Bateria", "Manguera PANI", "Manguito PANI", "Cable SPO2", 
        "Señales audibles", "Estado Pantalla/Display", "Estado Botones Interruptores", "Accesorios", 
        "Limpieza interior", "Calibración", "Cable ECG", "Estado de Palas", "Estado de Parches"
    ],
    "CALEFACCION": [
        "Chasis / Carcasa", "Limpieza exterior", "Cables alimentación", "Alarmas", 
        "Estado Bateria", "Señales audibles", "Estado Pantalla/Display", "Estado Botones Interruptores", "Accesorios", 
        "Limpieza interior"
    ],
    "CARPA": [
        "Recubrimiento", "Limpieza exterior", "Cables alimentación", "Accesorios", 
        "Limpieza interior", "Estado generador", "Mangueras", "Iluminarias", "Capas termicas"
    ],
    "CARROS": [
        "Chasis / Carcasa", "Limpieza exterior",
        "Limpieza interior", "Estado de ruedas", "Estado de frenos", "Lubricación", "Pernos y tuercas"
    ],
    "CHATAS": [
        "Chasis / Carcasa", "Limpieza exterior", "Cables alimentación", "Alarmas",  
        "Señales audibles", "Estado Pantalla/Display", "Estado Botones Interruptores", "Estado de sal",
        "Estado de detergente", "Sensores de nivel de Agua", "Estado Ablandador", "Parada de Emergencia"
    ],
    "ECOGRAFO": [
        "Chasis / Carcasa", "Limpieza exterior", "Cables alimentación", "Alarmas", 
        "Señales audibles", "Estado Pantalla/Display", "Estado Botones Interruptores", "Accesorios", 
        "Limpieza interior", "Estado de tranductores", "Limpieza de tranductores"
    ],
    "ELECTROVISTURI": [
        "Chasis / Carcasa", "Limpieza exterior", "Cables alimentación", "Alarmas", 
        "Estado Bateria", "Manguera PANI", "Manguito PANI", "Cable SPO2", 
        "Señales audibles", "Estado Pantalla/Display", "Estado Botones Interruptores", "Accesorios", 
        "Limpieza interior", "Calibración", "Cable ECG", "Estado de Palas", "Estado de Parches"
    ],
    "ESTERILIZACION": [
        "Chasis / Carcasa", "Limpieza exterior", "Cables alimentación", "Alarmas", 
        "Estado Bateria", "Manguera PANI", "Manguito PANI", "Cable SPO2", 
        "Señales audibles", "Estado Pantalla/Display", "Estado Botones Interruptores", "Accesorios", 
        "Limpieza interior", "Calibración", "Cable ECG", "Estado de Palas", "Estado de Parches"
    ],
    "GABINETE": [
        "Chasis / Carcasa", "Limpieza exterior", "Cables alimentación", "Alarmas", 
        "Señales audibles", "Accesorios", "Limpieza interior"
    ],
    "INCUBADORA": [
        "Chasis / Carcasa", "Limpieza exterior", "Cables alimentación", "Alarmas", 
        "Estado Bateria", "Cable SPO2", "Señales audibles", "Estado Pantalla/Display",
        "Estado Botones Interruptores", "Accesorios", 
        "Limpieza interior", "Calibración", "Estado de sellos", "Estado colochon"
    ],
    "MEDICIONES": [
        "Chasis / Carcasa", "Limpieza exterior", "Estado Pantalla/Display", 
        "Calibración", "Pilas / Baterias"
    ],
    "OPTICA": [
        "Chasis / Carcasa", "Limpieza exterior", "Cables alimentación", "Alarmas", 
        "Estado Bateria", "Manguera PANI", "Manguito PANI", "Cable SPO2", 
        "Señales audibles", "Estado Pantalla/Display", "Estado Botones Interruptores", "Accesorios", 
        "Limpieza interior", "Calibración", "Cable ECG", "Estado de Palas", "Estado de Parches"
    ],
    "PRESION": [
        "Chasis / Carcasa", "Limpieza exterior", "Cables alimentación", "Alarmas", 
        "Estado Bateria", "Manguera ", "Manguitos",
        "Señales audibles", "Estado Pantalla/Display", "Estado Botones Interruptores", "Accesorios",
        "Calibración"
    ],
    "REFRIGERACION": [
        "Chasis / Carcasa", "Limpieza exterior", "Cables alimentación", "Alarmas", 
        "Estado Bateria", "Señales audibles", "Estado Pantalla/Display",
        "Estado Botones Interruptores", "Limpieza interior", "Seteo temperaturas"
    ],
    "RX": [
        "Chasis / Carcasa", "Limpieza exterior", "Cables alimentación", "Alarmas", 
        "Estado Bateria", "Partes moviles", "Estado de ruedas",
        "Señales audibles", "Estado Pantalla/Display",
        "Estado Botones Interruptores", "Accesorios", "Calibración"
    ],
    "SELLADORA": [
        "Chasis / Carcasa", "Limpieza exterior", "Cables alimentación", "Alarmas", 
        "Señales audibles", "Estado Pantalla/Display", "Estado Botones Interruptores",
        "Limpieza interior", "Calibración"
    ],
    "TRASLADORAS": [
        "Chasis / Carcasa", "Limpieza exterior", 
        "Limpieza interior", "Estado de ruedas", "Estado de frenos",
        "Respaldo", "Estado colchon", "Estado Pistones", "Estado manillas", "Estado Barandas",
        "Estado de Bateria", "Controles / Perillas", "Indicadores / Display"
    ]
}

@app.get("/api/pautas")
def get_pautas():
    return PAUTAS_POR_CATEGORIA

# Servir uploads y archivos estáticos
app.mount("/uploads", StaticFiles(directory=UPLOADS_DIR), name="uploads")
app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")

@app.get("/")
def serve_home():
    return FileResponse(os.path.join(STATIC_DIR, "index.html"))

if __name__ == "__main__":
    import uvicorn
    print("Iniciando servidor EEMM en http://localhost:8000 ...")
    uvicorn.run("app:app", host="0.0.0.0", port=8000, reload=True)
